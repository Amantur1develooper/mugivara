"""Раздел «Отчёты» в личном кабинете.

Собирает заказы и записи по всем направлениям бизнеса за произвольный интервал
«дата+время → дата+время» и отдаёт либо таблицу на экран, либо файл Excel.
Доступ ограничен так же, как в аналитике: владелец видит только свои заведения,
суперадмин — все.
"""
from datetime import datetime, time, timedelta
from io import BytesIO

from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.shortcuts import render
from django.utils import timezone

LOGIN_URL = "dashboard:login"

# Сколько строк показываем на экране (в Excel уходят все)
PREVIEW_LIMIT = 300


def _sources():
    """Описание всех направлений: откуда брать записи и как называются поля.

    Ключи кортежа:
      model          — модель заказа/записи
      scope_field    — по какому полю ограничивать доступ пользователя
      ids_key        — ключ в словаре доступов пользователя
      venue_field    — путь к названию заведения
      customer/phone — поля клиента
      amount         — поле суммы (None — у направления нет суммы)
      related        — что подтянуть через select_related
    """
    from orders.models import Order
    from shops.models import StoreOrder
    from pharmacy.models import PharmacyOrder
    from hotels.models import HotelBooking
    from barbershop.models import Appointment as BarberAppointment
    from karaoke.models import KaraokeBooking
    from simracing.models import SimRacingAppointment
    from printshop.models import PrintOrder

    return [
        dict(key="rest", icon="🍽", label="Рестораны", model=Order,
             scope_field="branch__restaurant_id", ids_key="restaurant",
             venue_field="branch__restaurant__name_ru", related=("branch__restaurant",),
             customer="customer_name", phone="customer_phone", amount="total_amount"),
        dict(key="shop", icon="🏪", label="Магазины", model=StoreOrder,
             scope_field="branch__store_id", ids_key="store",
             venue_field="branch__store__name_ru", related=("branch__store",),
             customer="name", phone="phone", amount="total"),
        dict(key="pharm", icon="💊", label="Аптеки", model=PharmacyOrder,
             scope_field="branch__pharmacy_id", ids_key="pharmacy",
             venue_field="branch__pharmacy__name_ru", related=("branch__pharmacy",),
             customer="customer_name", phone="customer_phone", amount="total_amount"),
        dict(key="hotel", icon="🏨", label="Отели", model=HotelBooking,
             scope_field="branch__hotel_id", ids_key="hotel",
             venue_field="branch__hotel__name_ru", related=("branch__hotel",),
             customer="customer_name", phone="customer_phone", amount="total"),
        dict(key="barber", icon="✂️", label="Барбершопы", model=BarberAppointment,
             scope_field="barbershop_id", ids_key="barber",
             venue_field="barbershop__name", related=("barbershop",),
             customer="customer_name", phone="customer_phone", amount="price_snapshot"),
        dict(key="krk", icon="🎤", label="Караоке", model=KaraokeBooking,
             scope_field="venue_id", ids_key="karaoke",
             venue_field="venue__name", related=("venue",),
             customer="customer_name", phone="customer_phone", amount=None),
        dict(key="sim", icon="🏎️", label="Картинг", model=SimRacingAppointment,
             scope_field="venue_id", ids_key="simracing",
             venue_field="venue__name", related=("venue",),
             customer="customer_name", phone="customer_phone", amount="total_price"),
        dict(key="print", icon="🖨️", label="Полиграфия", model=PrintOrder,
             scope_field="branch__center_id", ids_key="printshop",
             venue_field="branch__center__name_ru", related=("branch__center",),
             customer="name", phone="phone", amount="total"),
    ]


def _user_ids(user):
    """ID заведений, к которым у пользователя есть доступ.
    None означает «без ограничений» (суперадмин)."""
    from core.models import Membership
    from shops.models import StoreMembership
    from pharmacy.models import PharmacyMembership
    from hotels.models import HotelMembership
    from barbershop.models import BarbershopMembership
    from karaoke.models import KaraokeMembership
    from simracing.models import SimRacingMembership
    from printshop.models import PrintMembership

    if user.is_superuser:
        return None

    def ids(model, field):
        return list(model.objects.filter(user=user).values_list(field, flat=True))

    return {
        "restaurant": ids(Membership, "restaurant_id"),
        "store":      ids(StoreMembership, "store_id"),
        "pharmacy":   ids(PharmacyMembership, "pharmacy_id"),
        "hotel":      ids(HotelMembership, "hotel_id"),
        "barber":     ids(BarbershopMembership, "barbershop_id"),
        "karaoke":    ids(KaraokeMembership, "venue_id"),
        "simracing":  ids(SimRacingMembership, "venue_id"),
        "printshop":  ids(PrintMembership, "center_id"),
    }


def _scoped_qs(src, user_ids):
    """Queryset направления, ограниченный доступом пользователя."""
    qs = src["model"].objects.all()
    if user_ids is None:
        return qs
    allowed = user_ids.get(src["ids_key"]) or []
    return qs.filter(**{f"{src['scope_field']}__in": allowed})


def venue_options(user):
    """Список заведений для выпадающего списка: [{value, label}, …].
    value — «направление:id», чтобы отчёт знал, где искать."""
    user_ids = _user_ids(user)
    options = []
    for src in _sources():
        qs = _scoped_qs(src, user_ids)
        rows = (qs.values(src["scope_field"], src["venue_field"])
                  .order_by(src["venue_field"]).distinct())
        for r in rows:
            vid = r[src["scope_field"]]
            name = r[src["venue_field"]]
            if not vid:
                continue
            options.append({
                "value": f"{src['key']}:{vid}",
                "label": f"{src['icon']} {name or '—'}",
            })
    return options


def parse_range(request):
    """Интервал отчёта из параметров запроса.
    По умолчанию — текущий месяц с начала и до конца сегодняшнего дня."""
    today = timezone.localtime(timezone.now()).date()

    def as_date(name, default):
        raw = (request.GET.get(name) or "").strip()
        try:
            return datetime.strptime(raw, "%Y-%m-%d").date()
        except ValueError:
            return default

    def as_time(name, default):
        raw = (request.GET.get(name) or "").strip()
        for fmt in ("%H:%M", "%H:%M:%S"):
            try:
                return datetime.strptime(raw, fmt).time()
            except ValueError:
                continue
        return default

    date_from = as_date("date_from", today.replace(day=1))
    date_to   = as_date("date_to", today)
    time_from = as_time("time_from", time(0, 0))
    time_to   = as_time("time_to", time(23, 59))

    if date_to < date_from:
        date_from, date_to = date_to, date_from

    naive_from = datetime.combine(date_from, time_from)
    naive_to   = datetime.combine(date_to, time_to)
    # конец интервала включительно: до конца указанной минуты
    naive_to += timedelta(minutes=1)

    return {
        "date_from": date_from, "date_to": date_to,
        "time_from": time_from, "time_to": time_to,
        "dt_from": timezone.make_aware(naive_from),
        "dt_to": timezone.make_aware(naive_to),
    }


def collect(user, venue_value, dt_from, dt_to):
    """Строки отчёта за интервал. venue_value — «направление:id» или пусто (все)."""
    user_ids = _user_ids(user)
    want_key, want_id = None, None
    if venue_value and ":" in venue_value:
        want_key, raw_id = venue_value.split(":", 1)
        if raw_id.isdigit():
            want_id = int(raw_id)

    rows, totals = [], {}
    for src in _sources():
        if want_key and src["key"] != want_key:
            continue
        qs = _scoped_qs(src, user_ids).filter(created_at__gte=dt_from, created_at__lt=dt_to)
        if want_id:
            qs = qs.filter(**{src["scope_field"]: want_id})
        qs = qs.select_related(*src["related"]).order_by("created_at")

        cnt = 0
        amount_sum = 0
        for obj in qs:
            created = timezone.localtime(obj.created_at)
            amount = 0
            if src["amount"]:
                amount = getattr(obj, src["amount"], 0) or 0
            status = (obj.get_status_display()
                      if hasattr(obj, "get_status_display") else getattr(obj, "status", ""))
            rows.append({
                "section": src["label"],
                "venue": _venue_name(obj, src["venue_field"]),
                "date": created.strftime("%d.%m.%Y"),
                "time": created.strftime("%H:%M"),
                "id": obj.id,
                "customer": getattr(obj, src["customer"], "") or "",
                "phone": getattr(obj, src["phone"], "") or "",
                "amount": amount,
                "status": status,
                "sort": created,
            })
            cnt += 1
            amount_sum += amount
        if cnt:
            totals[src["label"]] = {"icon": src["icon"], "count": cnt, "amount": amount_sum}

    rows.sort(key=lambda r: r["sort"])
    grand = {
        "count": sum(t["count"] for t in totals.values()),
        "amount": sum(t["amount"] for t in totals.values()),
    }
    return rows, totals, grand


def _venue_name(obj, path):
    """Название заведения по пути вида «branch__restaurant__name_ru»."""
    cur = obj
    for part in path.split("__"):
        cur = getattr(cur, part, None)
        if cur is None:
            return "—"
    return cur or "—"


# ── Views ────────────────────────────────────────────────────────────────────

@login_required(login_url=LOGIN_URL)
def reports(request):
    rng = parse_range(request)
    venue_value = (request.GET.get("venue") or "").strip()
    rows, totals, grand = collect(request.user, venue_value, rng["dt_from"], rng["dt_to"])

    export_qs = request.GET.urlencode()
    return render(request, "dashboard/reports.html", {
        "venues": venue_options(request.user),
        "venue_value": venue_value,
        "date_from": rng["date_from"].strftime("%Y-%m-%d"),
        "date_to": rng["date_to"].strftime("%Y-%m-%d"),
        "time_from": rng["time_from"].strftime("%H:%M"),
        "time_to": rng["time_to"].strftime("%H:%M"),
        "rows": rows[:PREVIEW_LIMIT],
        "rows_total": len(rows),
        "preview_limit": PREVIEW_LIMIT,
        "truncated": len(rows) > PREVIEW_LIMIT,
        "totals": sorted(totals.items(), key=lambda kv: -kv[1]["count"]),
        "grand": grand,
        "export_qs": export_qs,
    })


@login_required(login_url=LOGIN_URL)
def reports_export(request):
    """Тот же отчёт, но файлом Excel."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    rng = parse_range(request)
    venue_value = (request.GET.get("venue") or "").strip()
    rows, totals, grand = collect(request.user, venue_value, rng["dt_from"], rng["dt_to"])

    wb = Workbook()
    ws = wb.active
    ws.title = "Отчёт"

    accent = PatternFill("solid", fgColor="FF4D1C")
    head_font = Font(bold=True, color="FFFFFF", size=11)
    thin = Side(style="thin", color="DDDDDD")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    period = (f'{rng["date_from"]:%d.%m.%Y} {rng["time_from"]:%H:%M}'
              f' — {rng["date_to"]:%d.%m.%Y} {rng["time_to"]:%H:%M}')
    ws["A1"] = "Отчёт по заведению"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Период: {period}"
    ws["A3"] = f"Всего записей: {grand['count']}   Сумма: {int(grand['amount'])} сом"
    ws["A4"] = f"Сформирован: {timezone.localtime(timezone.now()):%d.%m.%Y %H:%M}"

    headers = ["Дата", "Время", "Направление", "Заведение", "№", "Клиент",
               "Телефон", "Сумма, сом", "Статус"]
    header_row = 6
    for col, name in enumerate(headers, start=1):
        c = ws.cell(row=header_row, column=col, value=name)
        c.fill, c.font, c.border = accent, head_font, border
        c.alignment = Alignment(horizontal="center", vertical="center")

    for i, r in enumerate(rows, start=header_row + 1):
        values = [r["date"], r["time"], r["section"], r["venue"], r["id"],
                  r["customer"], r["phone"], int(r["amount"]), r["status"]]
        for col, val in enumerate(values, start=1):
            c = ws.cell(row=i, column=col, value=val)
            c.border = border
            if col == 8:
                c.number_format = "#,##0"

    # Итоговая строка
    if rows:
        last = header_row + len(rows) + 1
        ws.cell(row=last, column=1, value="ИТОГО").font = Font(bold=True)
        ws.cell(row=last, column=5, value=grand["count"]).font = Font(bold=True)
        tc = ws.cell(row=last, column=8, value=int(grand["amount"]))
        tc.font = Font(bold=True)
        tc.number_format = "#,##0"

    widths = [12, 8, 14, 26, 8, 22, 18, 13, 16]
    for col, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    # Второй лист — сводка по направлениям
    if totals:
        ws2 = wb.create_sheet("Сводка")
        ws2["A1"] = "Направление"
        ws2["B1"] = "Записей"
        ws2["C1"] = "Сумма, сом"
        for col in "ABC":
            c = ws2[f"{col}1"]
            c.fill, c.font, c.border = accent, head_font, border
        for i, (label, t) in enumerate(
                sorted(totals.items(), key=lambda kv: -kv[1]["count"]), start=2):
            ws2.cell(row=i, column=1, value=label)
            ws2.cell(row=i, column=2, value=t["count"])
            c = ws2.cell(row=i, column=3, value=int(t["amount"]))
            c.number_format = "#,##0"
        ws2.column_dimensions["A"].width = 20
        ws2.column_dimensions["B"].width = 12
        ws2.column_dimensions["C"].width = 14

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    fname = f'webordo_otchet_{rng["date_from"]:%Y%m%d}_{rng["date_to"]:%Y%m%d}.xlsx'
    resp = HttpResponse(
        buf.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    resp["Content-Disposition"] = f'attachment; filename="{fname}"'
    return resp


# ── Отчёт по филиалу кафе: детализация до часов ───────────────────────────────

def _branch_or_403(request, branch_id):
    """Филиал с проверкой доступа. Возвращает (branch, None) либо (None, redirect)."""
    from django.shortcuts import get_object_or_404, redirect as _redirect
    from core.models import Branch, Membership
    branch = get_object_or_404(Branch.objects.select_related("restaurant"), id=branch_id)
    user = request.user
    allowed = (user.is_staff or user.is_superuser
               or Membership.objects.filter(user=user, restaurant=branch.restaurant).exists())
    if not allowed:
        return None, _redirect("dashboard:home")
    return branch, None


def _branch_data(branch, dt_from, dt_to):
    """Заказы филиала за интервал + разбивка по часам, типам и оплате."""
    from orders.models import Order

    qs = (Order.objects
          .filter(branch=branch, created_at__gte=dt_from, created_at__lt=dt_to)
          .select_related("table_place")
          .order_by("created_at"))

    type_labels = dict(Order.Type.choices)
    pay_labels = dict(Order.PaymentMethod.choices)

    rows = []
    by_hour = {}          # «13:00» → {count, amount}
    by_type = {}
    by_pay = {}
    revenue = 0
    cancelled = 0

    for o in qs:
        created = timezone.localtime(o.created_at)
        amount = o.total_amount or 0
        is_cancelled = o.status == Order.Status.CANCELLED
        if is_cancelled:
            cancelled += 1
        else:
            revenue += amount

        rows.append({
            "date": created.strftime("%d.%m.%Y"),
            "time": created.strftime("%H:%M"),
            "id": o.id,
            "type": type_labels.get(o.type, o.type),
            "customer": o.customer_name or "",
            "phone": o.customer_phone or "",
            "address": o.delivery_address or "",
            "amount": amount,
            "pay": pay_labels.get(o.payment_method, o.payment_method or ""),
            "status": o.get_status_display(),
            "cancelled": is_cancelled,
        })

        if not is_cancelled:
            key = created.strftime("%H:00")
            slot = by_hour.setdefault(key, {"count": 0, "amount": 0})
            slot["count"] += 1
            slot["amount"] += amount

            t = type_labels.get(o.type, o.type)
            bt = by_type.setdefault(t, {"count": 0, "amount": 0})
            bt["count"] += 1
            bt["amount"] += amount

            p = pay_labels.get(o.payment_method, o.payment_method or "—")
            bp = by_pay.setdefault(p, {"count": 0, "amount": 0})
            bp["count"] += 1
            bp["amount"] += amount

    hours = [{"hour": h, **by_hour[h]} for h in sorted(by_hour)]
    peak = max((h["amount"] for h in hours), default=0)
    for h in hours:
        h["bar"] = round(h["amount"] / peak * 100) if peak else 0

    paid_count = len(rows) - cancelled
    return {
        "rows": rows,
        "hours": hours,
        "by_type": sorted(by_type.items(), key=lambda kv: -kv[1]["amount"]),
        "by_pay": sorted(by_pay.items(), key=lambda kv: -kv[1]["amount"]),
        "count": len(rows),
        "cancelled": cancelled,
        "revenue": revenue,
        "avg": round(revenue / paid_count) if paid_count else 0,
    }


@login_required(login_url=LOGIN_URL)
def branch_report(request, branch_id):
    branch, deny = _branch_or_403(request, branch_id)
    if deny:
        return deny
    rng = parse_range(request)
    data = _branch_data(branch, rng["dt_from"], rng["dt_to"])
    return render(request, "dashboard/branch_report.html", {
        "branch": branch,
        "date_from": rng["date_from"].strftime("%Y-%m-%d"),
        "date_to": rng["date_to"].strftime("%Y-%m-%d"),
        "time_from": rng["time_from"].strftime("%H:%M"),
        "time_to": rng["time_to"].strftime("%H:%M"),
        "rows": data["rows"][:PREVIEW_LIMIT],
        "rows_total": data["count"],
        "preview_limit": PREVIEW_LIMIT,
        "truncated": data["count"] > PREVIEW_LIMIT,
        "hours": data["hours"],
        "by_type": data["by_type"],
        "by_pay": data["by_pay"],
        "count": data["count"],
        "cancelled": data["cancelled"],
        "revenue": data["revenue"],
        "avg": data["avg"],
        "export_qs": request.GET.urlencode(),
    })


@login_required(login_url=LOGIN_URL)
def branch_report_export(request, branch_id):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    branch, deny = _branch_or_403(request, branch_id)
    if deny:
        return deny
    rng = parse_range(request)
    d = _branch_data(branch, rng["dt_from"], rng["dt_to"])

    accent = PatternFill("solid", fgColor="FF4D1C")
    head_font = Font(bold=True, color="FFFFFF", size=11)
    thin = Side(style="thin", color="DDDDDD")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    bold = Font(bold=True)

    def head(ws, names, row=1):
        for col, name in enumerate(names, start=1):
            c = ws.cell(row=row, column=col, value=name)
            c.fill, c.font, c.border = accent, head_font, border
            c.alignment = Alignment(horizontal="center", vertical="center")

    period = (f'{rng["date_from"]:%d.%m.%Y} {rng["time_from"]:%H:%M}'
              f' — {rng["date_to"]:%d.%m.%Y} {rng["time_to"]:%H:%M}')

    wb = Workbook()

    # Лист 1 — по часам (главное, ради чего отчёт)
    ws = wb.active
    ws.title = "По часам"
    ws["A1"] = f"{branch.restaurant.name_ru} — {branch.name_ru}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Период: {period}"
    ws["A3"] = (f"Заказов: {d['count']}   Выручка: {int(d['revenue'])} сом   "
                f"Средний чек: {int(d['avg'])} сом   Отменённых: {d['cancelled']}")
    ws["A4"] = f"Сформирован: {timezone.localtime(timezone.now()):%d.%m.%Y %H:%M}"
    head(ws, ["Час", "Заказов", "Выручка, сом", "Средний чек, сом"], row=6)
    for i, h in enumerate(d["hours"], start=7):
        avg_h = round(h["amount"] / h["count"]) if h["count"] else 0
        for col, val in enumerate([h["hour"], h["count"], int(h["amount"]), int(avg_h)], start=1):
            c = ws.cell(row=i, column=col, value=val)
            c.border = border
            if col in (3, 4):
                c.number_format = "#,##0"
    if d["hours"]:
        last = 7 + len(d["hours"])
        ws.cell(row=last, column=1, value="ИТОГО").font = bold
        ws.cell(row=last, column=2, value=sum(h["count"] for h in d["hours"])).font = bold
        tc = ws.cell(row=last, column=3, value=int(d["revenue"]))
        tc.font, tc.number_format = bold, "#,##0"
    for col, w in enumerate([10, 12, 16, 18], start=1):
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = ws.cell(row=7, column=1)

    # Лист 2 — все заказы
    ws2 = wb.create_sheet("Заказы")
    head(ws2, ["Дата", "Время", "№", "Тип", "Клиент", "Телефон", "Адрес / стол",
               "Сумма, сом", "Оплата", "Статус"])
    for i, r in enumerate(d["rows"], start=2):
        values = [r["date"], r["time"], r["id"], r["type"], r["customer"], r["phone"],
                  r["address"], int(r["amount"]), r["pay"], r["status"]]
        for col, val in enumerate(values, start=1):
            c = ws2.cell(row=i, column=col, value=val)
            c.border = border
            if col == 8:
                c.number_format = "#,##0"
    for col, w in enumerate([12, 8, 8, 14, 20, 17, 28, 13, 13, 14], start=1):
        ws2.column_dimensions[get_column_letter(col)].width = w
    ws2.freeze_panes = ws2.cell(row=2, column=1)

    # Лист 3 — сводка по типам и оплате
    ws3 = wb.create_sheet("Сводка")
    head(ws3, ["Тип заказа", "Заказов", "Сумма, сом"])
    row = 2
    for label, t in d["by_type"]:
        ws3.cell(row=row, column=1, value=label)
        ws3.cell(row=row, column=2, value=t["count"])
        c = ws3.cell(row=row, column=3, value=int(t["amount"])); c.number_format = "#,##0"
        row += 1
    row += 1
    head(ws3, ["Способ оплаты", "Заказов", "Сумма, сом"], row=row)
    row += 1
    for label, p in d["by_pay"]:
        ws3.cell(row=row, column=1, value=label)
        ws3.cell(row=row, column=2, value=p["count"])
        c = ws3.cell(row=row, column=3, value=int(p["amount"])); c.number_format = "#,##0"
        row += 1
    for col, w in enumerate([22, 12, 14], start=1):
        ws3.column_dimensions[get_column_letter(col)].width = w

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = (f'otchet_{branch.id}_{rng["date_from"]:%Y%m%d}_{rng["date_to"]:%Y%m%d}.xlsx')
    resp = HttpResponse(
        buf.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    resp["Content-Disposition"] = f'attachment; filename="{fname}"'
    return resp
